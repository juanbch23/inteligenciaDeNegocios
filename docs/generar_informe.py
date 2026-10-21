"""Genera docs\\Informe_Proyecto_HM.docx: informe y guion de presentación del caso H&M.
Las capturas de los dashboards se toman de docs\\img\\ (d1.png, d2.png, d3.png) si existen.
Uso: python generar_informe.py"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AQUI = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(AQUI, "img")
SALIDA = os.path.join(AQUI, "Informe_Proyecto_HM.docx")
ROJO, OSCURO, GRIS = RGBColor(0xCC, 0x07, 0x1E), RGBColor(0x11, 0x18, 0x27), RGBColor(0x4B, 0x55, 0x63)

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.2); s.top_margin = s.bottom_margin = Cm(2)
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
for n, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
    doc.styles[n].font.name = "Calibri"; doc.styles[n].font.size = Pt(size); doc.styles[n].font.color.rgb = OSCURO

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcolor); tcPr.append(sh)

def p(text, bold=False, italic=False, color=None, size=None, align=None, space=4):
    par = doc.add_paragraph(); r = par.add_run(text); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = color
    if size: r.font.size = Pt(size)
    if align: par.alignment = align
    par.paragraph_format.space_after = Pt(space)
    return par

def rich(parts):
    """parts: lista de (texto, negrita)"""
    par = doc.add_paragraph()
    for t, b in parts: r = par.add_run(t); r.bold = b
    par.paragraph_format.space_after = Pt(4)
    return par

def bullets(items, style="List Bullet"):
    for it in items:
        par = doc.add_paragraph(style=style)
        if isinstance(it, tuple): r = par.add_run(it[0]); r.bold = True; par.add_run(it[1])
        else: par.add_run(it)
        par.paragraph_format.space_after = Pt(2)

def table(headers, rows, widths=None, head_color="111827"):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); r.font.size = Pt(9.5); shade(c, head_color)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def figura(nombre, pie):
    ruta = os.path.join(IMG, nombre)
    if os.path.exists(ruta):
        doc.add_picture(ruta, width=Cm(16.5)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p(f"[Captura pendiente: {nombre}]", italic=True, color=GRIS, align=WD_ALIGN_PARAGRAPH.CENTER)
    p(pie, italic=True, color=GRIS, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, space=10)

def storytelling(inicio, conflicto, evidencia, resolucion, accion):
    table(["Momento", "Mensaje"], [("Inicio", inicio), ("Conflicto", conflicto), ("Evidencia", evidencia),
                                   ("Resolución", resolucion), ("Acción", accion)], widths=[3, 13.5], head_color="CC071E")

# ======================================================================= PORTADA
p("INTELIGENCIA DE NEGOCIOS", bold=True, color=ROJO, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space=30)
p("H&M: recuperar la tienda sin regalar el margen", bold=True, color=OSCURO, size=24, align=WD_ALIGN_PARAGRAPH.CENTER)
p("Solución de inteligencia de negocios: ETL con SSIS, Data Warehouse, cubo tabular SSAS y dashboards en Power BI",
  color=GRIS, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space=30)
table(["Elemento", "Detalle"], [
    ("Datos", "31,788,324 transacciones · 1,371,980 clientes · 105,542 artículos · 20/09/2018 al 22/09/2020"),
    ("Arquitectura", "CSV / Excel → STAGE_HM → DWH_HM → cubo Tabular_HM → Power BI (conexión en directo)"),
    ("Herramientas", "SQL Server 2025 Developer, Integration Services, Analysis Services (tabular), Visual Studio 2022, Power BI Desktop"),
    ("Entregables", "Scripts SQL, solución H&M.sln (SSIS + tabular), informe HM_Dashboard.pbip, este documento"),
    ("Repositorio", "https://github.com/juanbch23/inteligenciaDeNegocios"),
], widths=[3.5, 13])
doc.add_page_break()

# ======================================================================= 1. RESUMEN
doc.add_heading("1. Resumen ejecutivo", 1)
p("A nivel global, H&M cumple su meta de ventas (110 % del Escenario 2). Sin embargo, frente al mismo periodo del año "
  "anterior vende 10.6 % menos: la tienda física cae 21.9 % y en abril de 2020 no vende nada por el cierre de la pandemia. "
  "El canal online sostiene el negocio, pero con cestas más pequeñas (−9.4 % de ticket), mientras la tienda liquida con cada "
  "vez más descuento (20.9 %) y se captan menos clientes nuevos.")
p("La solución integra los seis archivos fuente en un Data Warehouse en estrella, un cubo tabular con 73 medidas DAX y tres "
  "dashboards (estratégico, táctico y analítico) que responden tres preguntas: ¿vendemos por el canal correcto?, ¿a los "
  "clientes correctos? y ¿con el precio y las metas correctas? El hallazgo de mayor impacto: el cliente omnicanal compra "
  "3.8 veces más que el de un solo canal, y solo el 35.5 % de los clientes lo es.")
table(["Hallazgo", "Dato", "Decisión que habilita"], [
    ("La venta cae aunque se cumple la meta", "−10.6 % interanual; meta 110 %", "Separar efecto pandemia del desempeño comparable"),
    ("Metas mal repartidas", "Solo 16 de 56 secciones cumplen", "Recalcular metas por peso de cada sección"),
    ("Cestas online más pequeñas", "Ticket online −9.4 %", "Venta cruzada en el checkout"),
    ("La tienda liquida", "Descuento 17.1 % → 20.9 %", "Política de rebajas y traslado de stock"),
    ("El omnicanal vale más", "3.8 veces más venta por cliente", "Click & collect y cupones cruzados"),
    ("El club vale más", "+40 % por cliente; 33.8 % de penetración", "Activar a los socios inactivos"),
], widths=[5, 5, 6.5])

# ======================================================================= 2. CASO
doc.add_heading("2. El caso de negocio", 1)
doc.add_heading("2.1 Contexto", 2)
p("H&M vende moda por dos canales: tienda física (canal 1) y online (canal 2). En dos años registra 88,465 millones en "
  "ventas (unidad del campo price; la fuente no declara moneda), 1.36 millones de clientes compradores y 105 mil "
  "artículos agrupados en 56 secciones. El área comercial fija metas diarias por departamento/sección y temporada en tres "
  "escenarios (GoalsHyM.xlsx).")
doc.add_heading("2.2 Origen del problema", 2)
table(["Síntoma (1 ene – 22 sep)", "2019", "2020", "Cambio"], [
    ("Ventas totales", "33,348 M", "29,817 M", "−10.6 %"),
    ("Ventas tienda física", "8,316 M", "6,496 M", "−21.9 %"),
    ("Ventas online", "25,032 M", "23,321 M", "−6.8 %"),
    ("Participación online", "75.1 %", "78.2 %", "+3.1 pp"),
    ("Ticket por cesta online (año)", "12,724", "11,528", "−9.4 %"),
    ("Artículos por cesta online (año)", "4.25", "3.96", "−0.29"),
    ("Descuento implícito en tienda (año)", "19.4 %", "20.9 %", "+1.5 pp"),
    ("Clientes nuevos", "535,691 (año)", "245,404 (a sep.)", "menor captación"),
], widths=[6.5, 3.3, 3.3, 3.4])
p("En abril de 2020 la tienda física vendió 0 y en marzo cayó a 381 M (un mes normal ronda los 900 M). El canal online "
  "absorbió parte del golpe (abril 2020: +17 % frente a abril 2019), pero la caída empezó antes: enero y febrero de 2020 ya "
  "estaban 14 % y 8 % por debajo de 2019, y la temporada verano-2020 (diciembre 2019 a marzo 2020) es la única que no "
  "alcanza la meta.")
doc.add_heading("2.3 Preguntas de negocio", 2)
bullets([("¿Estamos vendiendo por el canal correcto? ", "tienda vs. online, omnicanalidad."),
         ("¿A los clientes correctos? ", "captación, recompra, retención y club."),
         ("¿Con el precio y las metas correctas? ", "descuento y cumplimiento por sección y temporada.")])

# ======================================================================= 3. ARQUITECTURA
doc.add_heading("3. Arquitectura de la solución", 1)
table(["Capa", "Componente", "Qué hace"], [
    ("Fuentes", "6 archivos (5 CSV + 1 Excel)", "Transacciones, clientes, artículos, calendario, temporadas y metas"),
    ("Staging", "STAGE_HM (esquema HM)", "Copia fiel de cada archivo; se vacía y recarga en cada ejecución"),
    ("Data Warehouse", "DWH_HM (estrella)", "6 dimensiones con comodín −1, 2 hechos, tablas de error y auditoría"),
    ("ETL", "SSIS: proyectos stage_H&M y DWH_H&M", "15 paquetes + 2 maestros con auditoría DIM_AUDIT"),
    ("Modelo semántico", "SSAS tabular Tabular_HM (compat. 1200)", "8 tablas, 9 relaciones, 73 medidas DAX"),
    ("Visualización", "Power BI HM_Dashboard.pbip", "3 dashboards con conexión en directo al cubo"),
], widths=[3, 5.5, 8])
doc.add_heading("3.1 Calidad de datos detectada y tratamiento", 2)
table(["Problema en la fuente", "Tratamiento"], [
    ("articles.csv y customers.csv terminan línea en LF; el resto en CRLF", "Delimitador de fila configurado por archivo en SSIS"),
    ("DimFecha.csv trae Fin_Semana = 0 en todos los días", "Recalculado con columna derivada (sábado y domingo = 1)"),
    ("FN y Active vienen como '1.0' o vacío", "Convertidos a 1/0 en DimCliente"),
    ("fashion_news_frequency mezcla NONE, None y vacío", "Normalizado a NONE / SIN DATO"),
    ("article_id con cero a la izquierda en artículos pero no en transacciones", "Ambos se cargan como entero"),
    ("Texto con codificación mixta (ChloÃ©)", "Se conserva como viene de la fuente"),
    ("Monto de GoalsHyM es una meta diaria", "Vista que expande la meta a cada día de su mes y temporada"),
], widths=[8, 8.5])

# ======================================================================= 4. ETL
doc.add_heading("4. ETL con SSIS y Data Warehouse", 1)
p("Cada paquete sigue la plantilla del curso: Task Prior RowCnt → Task Get AuditKey → TRUNCATE → flujo de datos → "
  "Task Post RowCnt → Update AuditDim. Los maestros ejecutan los paquetes dentro de contenedores de secuencia "
  "(Dimensiones y luego Hechos). Las dimensiones se cargan con procedimientos MERGE (SCD tipo 1) y los hechos con "
  "búsquedas (Lookup) y división condicional: las filas con alguna llave sin coincidencia van a la tabla de error.")
table(["Tabla", "Filas", "Validación"], [
    ("StageTransactions → FactVentas", "31,788,324", "Igual al archivo; 0 errores; montos con diferencia 0"),
    ("StageCustomers → DimCliente", "1,371,980", "Igual al archivo"),
    ("StageArticles → DimArticulo", "105,542", "Igual al archivo"),
    ("StageGoals → FactMetas", "27,738", "Igual al Excel; 0 errores"),
    ("StageFecha → DimFecha", "734", "210 fines de semana corregidos"),
    ("StageTemporada → DimTemporada", "9", "Igual al archivo"),
], widths=[6.5, 3, 7])
p("Tiempos de referencia: carga a STAGE ≈ 7 minutos (transacciones 6.8 min); carga al DWH ≈ 37 minutos; proceso del cubo "
  "≈ 5 minutos. Llaves foráneas activas y validadas. La auditoría queda en DWH_HM.dbo.DIM_AUDIT.", italic=True, color=GRIS)

# ======================================================================= 5. INDICADORES
doc.add_heading("5. Indicadores críticos", 1)
p("Las metas marcadas como “propuesta” no vienen en los datos: se basan en el desempeño histórico 2018-2019 y deben validarse "
  "con el área comercial. Todas las cifras se reproducen con sql/05_indicadores.sql.", italic=True, color=GRIS)
table(["#", "Indicador", "Fórmula", "Meta", "Valor", "Lectura"], [
    ("1", "Ventas e interanual", "Ventas / mismo periodo año ant. − 1", "≥ 0 % (prop.)", "−10.6 %", "Tienda −21.9 %, online −6.8 %"),
    ("2", "Cumplimiento de meta", "Ventas / meta diaria acumulada", "≥ 100 %", "110.3 %", "Solo 16 de 56 secciones cumplen"),
    ("3", "Participación online", "Ventas online / ventas", "Monitorear", "78.2 % (2020)", "100 % en abril 2020"),
    ("4", "Ticket por cesta", "Ventas / cestas", "≥ año ant. (prop.)", "9,643", "Online −9.4 %"),
    ("5", "Artículos por cesta", "Líneas / cestas", "≥ 4.0 online (prop.)", "3.46", "Online 4.25 → 3.96"),
    ("6", "Clientes nuevos", "1.ª compra en el periodo", "≥ 40 mil/mes (prop.)", "245,404 en 2020", "Captación a la baja"),
    ("7", "Recompra", "Clientes con compra en > 1 mes", "≥ 65 % (prop.)", "63.9 %", "Uno de cada tres compra una sola vez"),
    ("8", "Retención", "Compradores 2019 que vuelven en 2020", "≥ 65 % (prop.)", "60.8 %", "4 de cada 10 no volvieron"),
    ("9", "Club activo", "Venta/cliente activo vs. no", "Penetración ≥ 40 %", "+40 %; 33.8 %", "Vale más, pocos lo son"),
    ("10", "Omnicanal", "Clientes en ambos canales", "≥ 40 % (prop.)", "35.5 %; 3.8 veces", "Palanca de mayor impacto"),
    ("11", "Descuento implícito", "1 − precio / precio de lista", "≤ 17 % (prop.)", "Tienda 20.9 %", "La tienda liquida"),
    ("12", "Concentración", "Venta del top 20 % de clientes", "Monitorear", "69.6 %", "14 secciones = 80 % de la venta"),
], widths=[0.7, 3.2, 3.8, 2.6, 2.4, 3.8])
doc.add_heading("Reglas de cálculo", 3)
bullets(["Cesta = mismo cliente + mismo día + mismo canal (el dataset no trae número de pedido).",
         "Ticket = ventas / cestas; no se usa precio por línea como ticket.",
         "Precio de lista = precio máximo observado del artículo en el periodo (aproximación: no hay lista oficial).",
         "Meta = Escenario 2 como meta oficial; el monto del Excel es diario.",
         "Canal 1 = tienda física y canal 2 = online (interpretación habitual del dataset)."])
doc.add_heading("Del dato a la decisión", 3)
table(["KPI", "Caso detectado", "Decisión"], [
    ("Interanual", "Tienda −21.9 %", "Plan de reactivación de tiendas por temporada"),
    ("Cumplimiento", "35 secciones bajo el 50 %", "Metas proporcionales al peso histórico de cada sección"),
    ("Ticket y artículos por cesta", "Cestas online más pequeñas", "Recomendaciones y bundles en el checkout"),
    ("Captación y retención", "Menos nuevos, 60.8 % retención", "Primera compra online y reactivación segmentada"),
    ("Club y omnicanal", "+40 % y 3.8 veces", "Activar socios y convertir clientes a omnicanal"),
    ("Descuento", "Tienda 20.9 %", "Política de rebajas y traslado de stock al online"),
    ("Concentración", "Top 20 % = 69.6 %", "Programa de fidelización para clientes top"),
], widths=[4, 5, 7.5])

# ======================================================================= 6. DASHBOARDS
doc.add_heading("6. Los tres dashboards", 1)
p("Los tres dashboards se conectan en directo al cubo Tabular_HM. Comparten un estilo propio (encabezado oscuro con franja "
  "roja H&M, tarjetas blancas, rojo para destacar y gris para comparar) y una narrativa dinámica: un texto calculado con DAX "
  "que cambia con los filtros y resume la historia en una frase.")

doc.add_heading("6.1 Dashboard 1 — Estratégico: “¿Crecemos de forma sana?”", 2)
rich([("Audiencia: ", True), ("gerencia general, comercial y finanzas. ", False), ("Frecuencia: ", True), ("mensual.", False)])
bullets(["Seis KPIs con semáforo frente a la meta: ventas, interanual, cumplimiento, participación online, retención y descuento.",
         "Ventas mes a mes por año: muestra el hueco de marzo-abril 2020.",
         "Cascada: cuánto de la caída explica cada canal y grupo de producto.",
         "Mezcla de canal al 100 %: abril 2020 fue 100 % online.",
         "Captación de clientes nuevos por mes.",
         "Por defecto: año 2020 y Escenario 2; el gráfico mensual muestra un tooltip con la tendencia de 7 días.",
         "Nota de lectura: los primeros meses (sep-oct 2018) muestran muchos clientes nuevos porque los datos empiezan ahí; "
         "la comparación útil es entre 2019 y 2020.",
         "Nota de lectura: dentro de un solo año el cliente omnicanal compra 2.2 veces más; en todo el periodo, 3.8 veces."])
figura("d1.png", "Figura 1. Dashboard estratégico")
storytelling("Cumplimos la meta: 103 % en 2020 (110 % en todo el periodo).",
             "Pero vendemos 10.6 % menos que en el mismo periodo de 2019.",
             "La tienda cae 21.9 % (0 ventas en abril) y el online −6.8 %; ya caíamos en enero y febrero.",
             "El online sostiene el negocio y el cliente omnicanal compra 3.8 veces más.",
             "Reactivar la tienda como puerta de entrada al omnicanal y separar el efecto pandemia en las metas 2021.")

doc.add_heading("6.2 Dashboard 2 — Táctico: “Torre de control de metas”", 2)
rich([("Audiencia: ", True), ("jefes de sección y de temporada. ", False), ("Frecuencia: ", True), ("diaria / por temporada.", False)])
bullets(["KPIs: ventas, meta, cumplimiento con semáforo, secciones que cumplen y secciones bajo el 50 %.",
         "Mapa de calor sección × temporada: verde ≥ 100 %, ámbar ≥ 90 %, rojo < 90 %.",
         "Dispersión meta vs. venta por sección: las burbujas lejos de la diagonal son metas mal asignadas.",
         "Top 5 y bottom 5 de cumplimiento.",
         "Curva de venta acumulada vs. meta acumulada para seguir una temporada día a día."])
figura("d2.png", "Figura 2. Dashboard táctico")
storytelling("La meta global se cumple (110 %).",
             "Solo 16 de 56 secciones la alcanzan y 35 no llegan ni a la mitad.",
             "Las metas son casi iguales para todas las secciones: Womens Swimwear llega a 1,918 % y Ladies Other a 1 %.",
             "Las metas deben reflejar el peso histórico de cada sección.",
             "Recalcular las metas por sección para la próxima temporada y seguirlas con la curva acumulada.")

doc.add_heading("6.3 Dashboard 3 — Analítico: “Laboratorio de cliente y precio”", 2)
rich([("Audiencia: ", True), ("marketing, pricing y BI. ", False), ("Frecuencia: ", True), ("semanal / mensual.", False)])
bullets(["KPIs: ticket y artículos por cesta, recompra, % omnicanal, multiplicador omnicanal y venta del top 20 % (sin filtro de año: todo el periodo).",
         "Árbol de descomposición: canal → grupo → sección → tipo de cliente → edad (en la presentación se abre en vivo con el botón +).",
         "Venta por cliente según tipo de cliente y según club activo por edad.",
         "Descuento implícito por canal y líneas por banda de descuento.",
         "Ticket y artículos por cesta por mes; Pareto de clientes por segmento de valor."])
figura("d3.png", "Figura 3. Dashboard analítico")
storytelling("1.36 millones de clientes compran en H&M.",
             "Cada vez compran menos por visita y se captan menos clientes nuevos.",
             "Ticket online −9.4 %; el cliente omnicanal vale 3.8 veces más y el de club activo 40 % más; la tienda vende con 20.9 % de descuento.",
             "Convertir clientes de un canal en omnicanal y activar el club, sin depender de la rebaja.",
             "Campañas segmentadas por edad y canal, click & collect y política de rebajas por temporada.")

# ======================================================================= 7. RECOMENDACIONES
doc.add_heading("7. Recomendaciones", 1)
bullets([("Omnicanal primero: ", "click & collect, devoluciones en tienda de compras online y cupones cruzados. Subir de 35.5 % a 40 % de clientes omnicanal es la palanca de mayor impacto."),
         ("Metas por sección: ", "recalcularlas con base en la participación histórica; hoy 35 secciones no llegan al 50 %."),
         ("Tienda sin liquidar: ", "política de rebajas por temporada y traslado de stock al canal online para bajar el descuento de 20.9 % a ≤ 17 %."),
         ("Cestas más grandes: ", "recomendaciones 'completa el look' y umbral de envío gratis por encima del ticket actual."),
         ("Clientes: ", "reactivar al 39 % que no volvió en 2020 y activar a los 808 mil socios ACTIVE sin Active = 1.")])

# ======================================================================= 8. GUION
doc.add_heading("8. Guion de la presentación (15 minutos)", 1)
table(["#", "Tiempo", "Tema", "Qué mostrar / decir"], [
    ("1", "1 min", "Apertura", "El titular: cumplimos la meta pero vendemos 10.6 % menos. Las tres preguntas de negocio."),
    ("2", "2 min", "Datos y arquitectura", "6 fuentes, 31.8 M de transacciones; CSV → STAGE → DWH → cubo → Power BI."),
    ("3", "2 min", "ETL y calidad", "Paquetes con auditoría, búsquedas y tabla de error; problemas de calidad encontrados y cómo se resolvieron."),
    ("4", "1 min", "Modelo y cubo", "Estrella con 6 dimensiones y 2 hechos; 73 medidas DAX; meta diaria."),
    ("5", "3 min", "Dashboard estratégico", "Recorrer la narrativa: KPIs en rojo, hueco de abril 2020, cascada por canal."),
    ("6", "2 min", "Dashboard táctico", "Mapa de calor y dispersión: metas mal asignadas; filtrar una temporada y ver la curva acumulada."),
    ("7", "3 min", "Dashboard analítico", "Árbol de descomposición en vivo; omnicanal 3.8 veces; descuento de la tienda."),
    ("8", "1 min", "Cierre", "Cinco recomendaciones y la decisión que habilita cada indicador."),
], widths=[0.7, 1.6, 3.6, 10.6])
doc.add_heading("Preguntas probables y respuesta", 3)
table(["Pregunta", "Respuesta"], [
    ("¿Por qué la meta es diaria?", "Como meta mensual el cumplimiento daría ~2.800 %; como diaria da 127/110/95 % en los tres escenarios, que es coherente."),
    ("¿Cómo calculan el ticket sin número de pedido?", "Cesta = mismo cliente, día y canal. Es la convención habitual con este dataset."),
    ("¿De dónde sale el descuento?", "Precio cobrado vs. el precio máximo observado del mismo artículo; es una aproximación al precio de lista."),
    ("¿Cómo saben que canal 2 es online?", "Es la interpretación habitual del dataset; además, en abril 2020 (tiendas cerradas) solo vendió el canal 2."),
    ("¿Cómo validaron las cifras?", "03_validaciones.sql (archivo vs. STAGE vs. DWH, montos con diferencia 0) y 05_indicadores.sql contra las medidas DAX."),
], widths=[5.5, 11])

# ======================================================================= 9. LIMITACIONES
doc.add_heading("9. Supuestos y limitaciones", 1)
bullets(["Los datos terminan el 22/09/2020: 2018 y 2020 son años incompletos; por eso se compara el mismo periodo.",
         "La fuente no declara moneda; los importes están en la unidad del campo price.",
         "El precio de lista es aproximado (máximo observado); no hay costos, por lo que no se calcula margen.",
         "Las metas distintas a ventas son propuestas y deben validarse con el negocio.",
         "Los clientes “nuevos” de 2018-2019 incluyen a quienes ya compraban antes del 20/09/2018 (inicio de los datos)."])

doc.save(SALIDA)
print("OK", SALIDA)
